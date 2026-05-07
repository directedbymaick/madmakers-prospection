"""
enrich_dirigeants_linkedin.py — Enrichit le batch_france50_active avec :
  1. Dirigeant via societe.com (Pappers epuise) :
     - /manager/Prenom.NOM.<hash>.html (cas standard)
     - <p class="is-Fit">{Role} : NOM, Prenom</p> (fallback pour entreprises recentes)
  2. URL LinkedIn via Bing search ("site:linkedin.com/in <Prenom Nom> <Societe>")

Sortie : nouveau JSON + Excel avec colonnes Dirigeant / Fonction / LinkedIn.
"""
import json, time, re, requests, sys
from pathlib import Path
from urllib.parse import quote_plus
from datetime import datetime

DATA_DIR = Path(r"C:\Users\MAÏCK\Desktop\MadMakers Prospection\linkedin-scraper\data")
TODAY = datetime.now().strftime("%Y%m%d")

INPUT_JSON  = DATA_DIR / f"batch_france50_active_{TODAY}.json"
OUTPUT_JSON = DATA_DIR / f"batch_france50_active_enriched_{TODAY}.json"

HDR = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    "Accept-Language": "fr-FR,fr;q=0.9",
}

ROLES_RE = r'(?:Pr[ée]sident|G[ée]rant|Co-?G[ée]rant|Directeur g[ée]n[ée]ral|PDG|Administrateur)'
PAT_MGR = re.compile(r'/manager/([\w\-]+)\.([\w\-]+)\.\w+\.html')
PAT_FIT = re.compile(
    r'<p class="is-Fit">\s*(' + ROLES_RE + r')\s*:\s*'
    r'([A-Z][^<,]{1,40}?),\s*([^<,]{1,40}?)\s*</p>'
)


def societe_dirigeant(siren: str) -> dict:
    """Cherche le dirigeant sur societe.com par SIREN. Retourne {prenom, nom, fonction}."""
    url = f"https://www.societe.com/cgi-bin/recherche?rncs={siren}"
    try:
        r = requests.get(url, headers=HDR, timeout=15, allow_redirects=True)
    except requests.RequestException as e:
        return {"_error": f"net_{type(e).__name__}"}
    if r.status_code == 404:
        return {"_error": "404"}
    if r.status_code != 200:
        return {"_error": f"http_{r.status_code}"}

    # Essai 1 : lien /manager/Prenom.NOM.<hash>.html
    m = PAT_MGR.search(r.text)
    if m:
        prenom = m.group(1).replace(".", " ").strip()
        nom    = m.group(2).strip()
        return {"dirigeant_prenom": prenom, "dirigeant_nom": nom,
                "dirigeant_fonction": "", "_source": "manager_link"}

    # Essai 2 : <p class="is-Fit">Role : NOM, Prenom</p>
    m = PAT_FIT.search(r.text)
    if m:
        role  = m.group(1).strip()
        nom   = m.group(2).strip()
        prenom= m.group(3).strip()
        return {"dirigeant_prenom": prenom, "dirigeant_nom": nom,
                "dirigeant_fonction": role, "_source": "is_fit"}

    return {}


def bing_linkedin(prenom: str, nom: str, company: str, ville: str) -> str:
    """Cherche profil LinkedIn via Bing. Retourne URL ou ''."""
    full_name = f"{prenom} {nom}".strip()
    if not full_name:
        return ""

    queries = [
        f'site:linkedin.com/in "{full_name}" "{company}"',
        f'site:linkedin.com/in "{full_name}" "{ville}"',
    ]
    for q in queries:
        url = f"https://www.bing.com/search?q={quote_plus(q)}&mkt=fr-FR&setlang=fr"
        try:
            r = requests.get(url, headers=HDR, timeout=12)
        except requests.RequestException:
            continue
        if r.status_code != 200:
            continue
        li_urls = re.findall(
            r'href="(https://(?:www\.)?linkedin\.com/in/[^"?&]{3,80})"',
            r.text)
        seen, clean = set(), []
        for u in li_urls:
            u2 = u.split("?")[0].rstrip("/")
            if u2 not in seen:
                seen.add(u2)
                clean.append(u2)
        if clean:
            return clean[0]
        time.sleep(1.0)
    return ""


def main():
    if not INPUT_JSON.exists():
        print(f"ERREUR : {INPUT_JSON} introuvable")
        sys.exit(1)
    with open(INPUT_JSON, encoding="utf-8") as f:
        prospects = json.load(f)

    print(f"Enrichissement {len(prospects)} prospects (societe.com + Bing)\n")

    n_dirigeant = 0
    n_linkedin  = 0
    sources = {"manager_link": 0, "is_fit": 0}
    errors = []

    for i, p in enumerate(prospects, 1):
        nom_resto = p['nom_commercial'][:30]
        print(f"[{i:>2}/{len(prospects)}] {nom_resto:30s}", end=" | ", flush=True)

        d = societe_dirigeant(p['siren'])
        if "_error" in d:
            errors.append((p['siren'], d['_error']))
            p['dirigeant_prenom']   = ""
            p['dirigeant_nom']      = ""
            p['dirigeant_fonction'] = ""
            p['linkedin_url']       = ""
            print(f"societe.com {d['_error']}")
            time.sleep(0.4)
            continue

        if d.get("dirigeant_nom"):
            p['dirigeant_prenom']   = d['dirigeant_prenom']
            p['dirigeant_nom']      = d['dirigeant_nom']
            p['dirigeant_fonction'] = d.get('dirigeant_fonction', '')
            n_dirigeant += 1
            sources[d['_source']] = sources.get(d['_source'], 0) + 1
            full = f"{d['dirigeant_prenom']} {d['dirigeant_nom']}"
            print(f"{full[:26]:26s}", end=" | ", flush=True)

            time.sleep(0.5)
            li = bing_linkedin(d['dirigeant_prenom'], d['dirigeant_nom'],
                               p['nom_commercial'], p['ville'])
            p['linkedin_url'] = li
            if li:
                n_linkedin += 1
                tail = li.replace('https://www.linkedin.com/in/', '').replace('https://linkedin.com/in/', '')
                print(f"LI {tail[:30]}")
            else:
                print("LI -")
        else:
            p['dirigeant_prenom']   = ""
            p['dirigeant_nom']      = ""
            p['dirigeant_fonction'] = ""
            p['linkedin_url']       = ""
            print("(aucun dirigeant trouve)")

        time.sleep(0.4)

    print(f"\n=== Resultats ===")
    print(f"Dirigeants trouves : {n_dirigeant}/{len(prospects)}  (sources: {sources})")
    print(f"LinkedIn trouves   : {n_linkedin}/{len(prospects)}")
    if errors:
        from collections import Counter
        print(f"Erreurs societe.com: {dict(Counter(e[1] for e in errors))}")

    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(prospects, f, ensure_ascii=False, indent=2)
    print(f"\nJSON enrichi : {OUTPUT_JSON}")

    make_excel_enriched(prospects)


def make_excel_enriched(prospects):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Prospects enrichis"

    hdr_fill = PatternFill("solid", fgColor="1a3c6e")
    hdr_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    alt_fill = PatternFill("solid", fgColor="EEF2FA")
    wht_fill = PatternFill("solid", fgColor="FFFFFF")
    grn_fill = PatternFill("solid", fgColor="C8E6C9")
    org_fill = PatternFill("solid", fgColor="FFE0B2")
    ylw_fill = PatternFill("solid", fgColor="FFF9C4")
    thin     = Side(style="thin", color="CCCCCC")
    border   = Border(left=thin, right=thin, top=thin, bottom=thin)
    center   = Alignment(horizontal="center", vertical="center")
    left_al  = Alignment(horizontal="left", vertical="center")

    headers = ["#", "Nom commercial", "Raison sociale", "SIRET", "SIREN",
               "Secteur", "Adresse", "CP", "Ville", "Departement",
               "Telephone", "Note Google", "Nb avis",
               "Dirigeant", "Fonction", "LinkedIn", "Email", "Statut"]
    col_w = [4, 28, 26, 16, 12, 22, 28, 7, 18, 22,
             14, 10, 8, 26, 18, 52, 26, 13]

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hdr_font; c.fill = hdr_fill
        c.alignment = center; c.border = border

    for i, p in enumerate(prospects):
        r = i + 2
        base = alt_fill if r % 2 == 1 else wht_fill
        rating = p.get("note_google")
        prenom = p.get("dirigeant_prenom","") or ""
        nom    = p.get("dirigeant_nom","") or ""
        dirigeant_full = f"{prenom} {nom}".strip()

        vals = [p["num"], p["nom_commercial"], p["raison_sociale"],
                p["siret"], p["siren"], p["secteur"],
                p["adresse"], p["cp"], p["ville"], p["departement"],
                p["telephone"], rating, p["nb_avis"],
                dirigeant_full, p.get("dirigeant_fonction","") or "",
                p.get("linkedin_url","") or "",
                p.get("email","") or "",
                p.get("statut","A contacter")]

        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.border = border
            cell.font = Font(name="Arial", size=9)
            cell.alignment = center if col in [1,4,5,8,12,13] else left_al

            if col == 16 and str(val).startswith("http"):
                cell.hyperlink = str(val)
                cell.font = Font(name="Arial", size=9, color="0563C1", underline="single")
                cell.fill = grn_fill
            elif col == 12 and isinstance(rating, float) and rating >= 4.5:
                cell.fill = grn_fill
            elif col == 14 and val:
                cell.fill = grn_fill
            elif col == 17:
                cell.fill = org_fill
            elif col == 18:
                cell.fill = ylw_fill
            else:
                cell.fill = base

    for i, w in enumerate(col_w, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    for r in range(2, len(prospects) + 2):
        ws.row_dimensions[r].height = 18
    ws.freeze_panes = "A2"
    if prospects:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(prospects)+1}"

    out = DATA_DIR / f"prospects_france50_active_enriched_{TODAY}.xlsx"
    wb.save(out)
    print(f"Excel : {out}")


if __name__ == "__main__":
    main()
