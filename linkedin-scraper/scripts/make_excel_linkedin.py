import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

df_raw = pd.read_excel('C:/Users/MAÏCK/Desktop/MadMakers Prospection/linkedin-scraper/data/prospects_sans_site.xlsx')

df = df_raw[df_raw['Signal'].isin(['sans_site','sans_https'])].copy()

exclude_sectors = ['Organisations civiques et sociales','Administration publique',
    'Services gouvernementaux','Organisations religieuses','Enseignement primaire et secondaire']
exclude_names = ['CCI','Association','Fédération','CITOYENS ET JUSTICE']
df = df[~df['Secteur'].isin(exclude_sectors)]
for excl in exclude_names:
    df = df[~df['Entreprise'].str.contains(excl, case=False, na=False)]

def score(row):
    s = row['Score']
    loc = str(row['Localisation'])
    if 'Nouvelle-Aquitaine' in loc: s += 20
    if any(x in loc for x in ['Bordeaux','Landes','Pays Basque','Bayonne','Pau','Périgueux']): s += 10
    return s

df['Score'] = df.apply(score, axis=1)
df = df.sort_values('Score', ascending=False).reset_index(drop=True)

wb = Workbook()
ws = wb.active
ws.title = "Prospects LinkedIn"

header_fill = PatternFill("solid", fgColor="1a3c6e")
header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
alt_fill    = PatternFill("solid", fgColor="EEF2FA")
white_fill  = PatternFill("solid", fgColor="FFFFFF")
red_fill    = PatternFill("solid", fgColor="FFCCCC")
yellow_fill = PatternFill("solid", fgColor="FFF3CD")
orange_fill = PatternFill("solid", fgColor="FFE0B2")
thin        = Side(style="thin", color="CCCCCC")
border      = Border(left=thin, right=thin, top=thin, bottom=thin)
center      = Alignment(horizontal="center", vertical="center")
left        = Alignment(horizontal="left", vertical="center", wrap_text=True)

headers = ["#","Nom","Titre","Entreprise","Secteur","Localisation","Signal","Score","LinkedIn","Statut"]
col_widths = [4, 22, 28, 30, 25, 30, 12, 8, 45, 14]

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center
    cell.border = border

for i, row in df.iterrows():
    r = i + 2
    is_alt = (r % 2 == 1)
    base = alt_fill if is_alt else white_fill
    signal = str(row.get('Signal',''))

    vals = [
        i + 1,
        str(row.get('Nom','')),
        str(row.get('Titre','')),
        str(row.get('Entreprise','')),
        str(row.get('Secteur','')) if pd.notna(row.get('Secteur')) else '',
        str(row.get('Localisation','')) if pd.notna(row.get('Localisation')) else '',
        signal,
        int(row.get('Score', 0)),
        str(row.get('LinkedIn','')),
        'A contacter',
    ]

    for col, val in enumerate(vals, 1):
        cell = ws.cell(row=r, column=col, value=val)
        cell.border = border
        cell.font = Font(name="Arial", size=9)
        cell.alignment = center if col in [1, 8] else left

        if col == 7:  # Signal
            cell.fill = red_fill if signal == 'sans_site' else yellow_fill
        elif col == 9:  # LinkedIn
            if val and val.startswith('http'):
                cell.hyperlink = val
                cell.font = Font(name="Arial", size=9, color="0563C1", underline="single")
                cell.fill = base
            else:
                cell.fill = base
        elif col == 10:  # Statut
            cell.fill = orange_fill
        else:
            cell.fill = base

for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.row_dimensions[1].height = 30
for r in range(2, len(df) + 2):
    ws.row_dimensions[r].height = 18

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(df)+1}"

out = "C:/Users/MAÏCK/Desktop/MadMakers Prospection/linkedin-scraper/data/madmakers_linkedin_prospects.xlsx"
wb.save(out)
print(f"✅ {len(df)} prospects — {out}")
